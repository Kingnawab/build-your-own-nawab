from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from duty_swap_automation.config import Settings
from duty_swap_automation.mailbox import LocalMailbox
from duty_swap_automation.main import DutySwapFlow
from duty_swap_automation.models import ParsedEmail, SwapRequest
from duty_swap_automation.roster import LocalWorkbook, SwapError, load_roster, plan_changes, to_date
from duty_swap_automation.sample_data import build

MONDAY = date.today() + timedelta(days=(7 - date.today().weekday()))


@pytest.fixture
def settings(tmp_path):
    build(tmp_path, MONDAY)
    return Settings(
        backend="local",
        timezone="UTC",
        local_workbook_path=str(tmp_path / "duty_roster.xlsx"),
        local_mail_dir=str(tmp_path / "mail"),
    )


def assignee(settings, day, duty):
    ws = load_workbook(settings.local_workbook_path)["Roster"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        if to_date(row[0]) == day and row[2] == duty:
            return row[3]
    raise AssertionError("row not found")


def who_has(settings, day):
    roster = load_roster(LocalWorkbook(settings), settings)
    return {r.duty: r.assignee for r in roster.rows if r.date == day}


def stub_parser(mapping):
    def parse(email, roster, settings):
        return mapping.get(email.id, ParsedEmail(is_swap_request=False))
    return parse


def test_swap_updates_roster_logs_and_replies(settings):
    tue, wed = MONDAY + timedelta(days=1), MONDAY + timedelta(days=2)
    tue_gate = who_has(settings, tue)["Morning Gate"]
    wed_gate = who_has(settings, wed)["Morning Gate"]
    req = SwapRequest(request_type="swap", person_a=wed_gate, date_a=wed, duty_a="Morning Gate",
                      person_b=tue_gate, date_b=tue, duty_b="Morning Gate", summary="swap")
    parsed = {"msg-001": ParsedEmail(is_swap_request=True, requests=[req])}

    settings.require_requester_is_party = False
    flow = DutySwapFlow(settings=settings, parse_fn=stub_parser(parsed))
    flow.kickoff()

    assert assignee(settings, tue, "Morning Gate") == wed_gate
    assert assignee(settings, wed, "Morning Gate") == tue_gate
    assert [o.status for o in flow.state.outcomes] == ["applied"]

    log_rows = list(load_workbook(settings.local_workbook_path)["Swap Log"].values)
    assert len(log_rows) == 2

    mailbox = LocalMailbox(settings)
    reply = (mailbox.root / "outbox" / "reply-msg-001.txt").read_text()
    assert "have been completed" in reply
    assert "msg-001" in mailbox._state() and "msg-002" not in mailbox._state()
    assert [e.id for e in mailbox.fetch_candidates()] == ["msg-002"]


def test_unauthorized_requester_is_rejected(settings):
    tue, wed = MONDAY + timedelta(days=1), MONDAY + timedelta(days=2)
    tue_gate = who_has(settings, tue)["Morning Gate"]
    wed_gate = who_has(settings, wed)["Morning Gate"]
    req = SwapRequest(request_type="swap", person_a=wed_gate, date_a=wed, duty_a="Morning Gate",
                      person_b=tue_gate, date_b=tue, duty_b="Morning Gate")
    path = LocalMailbox(settings).root / "inbox" / "msg-002.json"
    path.write_text(path.read_text().replace("Chen Wei", "Mallory Outsider").replace("chen@", "mallory@"))
    flow = DutySwapFlow(settings=settings, parse_fn=stub_parser({"msg-002": ParsedEmail(is_swap_request=True, requests=[req])}))
    flow.kickoff()
    assert [o.status for o in flow.state.outcomes] == ["rejected"]
    assert assignee(settings, tue, "Morning Gate") == tue_gate


def test_dry_run_changes_nothing(settings):
    tue, wed = MONDAY + timedelta(days=1), MONDAY + timedelta(days=2)
    tue_gate = who_has(settings, tue)["Morning Gate"]
    wed_gate = who_has(settings, wed)["Morning Gate"]
    req = SwapRequest(request_type="swap", person_a=wed_gate, date_a=wed, duty_a="Morning Gate",
                      person_b=tue_gate, date_b=tue, duty_b="Morning Gate")
    settings.dry_run = True
    settings.require_requester_is_party = False
    flow = DutySwapFlow(settings=settings, parse_fn=stub_parser({"msg-001": ParsedEmail(is_swap_request=True, requests=[req])}))
    flow.kickoff()
    assert [o.status for o in flow.state.outcomes] == ["dry_run"]
    assert assignee(settings, tue, "Morning Gate") == tue_gate
    assert not (LocalMailbox(settings).root / "outbox").exists()


def test_clarification_reply(settings):
    parsed = {"msg-001": ParsedEmail(is_swap_request=True, needs_clarification=True,
                                     clarification_question="Which day would you like to swap?")}
    flow = DutySwapFlow(settings=settings, parse_fn=stub_parser(parsed))
    flow.kickoff()
    reply = (LocalMailbox(settings).root / "outbox" / "reply-msg-001.txt").read_text()
    assert "Which day would you like to swap?" in reply
    assert LocalMailbox(settings)._state()["msg-001"] == settings.attention_category


def test_plan_rejects_wrong_person_and_past_dates(settings):
    roster = load_roster(LocalWorkbook(settings), settings)
    tue = MONDAY + timedelta(days=1)
    gate = who_has(settings, tue)["Morning Gate"]
    other = next(p for p in roster.people if p not in who_has(settings, tue).values())
    with pytest.raises(SwapError, match="not on the roster"):
        plan_changes(roster, SwapRequest(request_type="cover", person_a=other, date_a=tue, person_b=gate),
                     date.today(), False)
    with pytest.raises(SwapError, match="already passed"):
        plan_changes(roster, SwapRequest(request_type="cover", person_a=gate, date_a=tue, person_b=other),
                     tue + timedelta(days=1), False)
    changes, _ = plan_changes(roster, SwapRequest(request_type="cover", person_a=gate, date_a=tue, duty_a="gate",
                                                  person_b=other), date.today(), False)
    assert len(changes) == 1 and changes[0].new_value == other


def test_excel_serial_dates_from_graph():
    assert to_date(46000) == date(2025, 12, 9)
    assert to_date("2026-11-03") == date(2026, 11, 3)
