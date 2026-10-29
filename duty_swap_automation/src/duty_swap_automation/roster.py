from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Protocol
from urllib.parse import quote

from dateutil import parser as date_parser
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from duty_swap_automation.config import Settings
from duty_swap_automation.graph_client import GraphClient, GraphError
from duty_swap_automation.models import CellChange, SwapRequest

EXCEL_EPOCH = date(1899, 12, 30)


class SwapError(ValueError):
    pass


@dataclass
class SheetData:
    values: list[list[Any]]
    top_row: int = 1  # 1-based worksheet row of values[0]
    left_col: int = 1  # 1-based worksheet column of values[i][0]


class Workbook(Protocol):
    def read_sheet(self, sheet: str) -> SheetData | None: ...

    def write_cells(self, changes: list[CellChange]) -> None: ...

    def append_row(self, sheet: str, header: list[str], row: list[Any]) -> None: ...

    def close(self) -> None: ...


def _a1(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


class GraphWorkbook:
    def __init__(self, client: GraphClient, settings: Settings):
        self.client = client
        s = settings
        if s.roster_drive_id and s.roster_item_id:
            self.base = f"/drives/{s.roster_drive_id}/items/{s.roster_item_id}/workbook"
        elif s.roster_site_id and s.roster_file_path:
            self.base = f"/sites/{s.roster_site_id}/drive/root:/{quote(s.roster_file_path.strip('/'))}:/workbook"
        elif s.roster_file_path:
            self.base = f"{client.user_base}/drive/root:/{quote(s.roster_file_path.strip('/'))}:/workbook"
        else:
            raise GraphError("Set ROSTER_FILE_PATH (optionally with ROSTER_SITE_ID) or ROSTER_DRIVE_ID + ROSTER_ITEM_ID.")
        self._session_id: str | None = None

    def _headers(self) -> dict:
        if self._session_id is None:
            resp = self.client.request("POST", f"{self.base}/createSession", json={"persistChanges": True})
            self._session_id = resp["id"]
        return {"workbook-session-id": self._session_id}

    def _sheet_url(self, sheet: str) -> str:
        return f"{self.base}/worksheets('{quote(sheet.replace(chr(39), chr(39) * 2))}')"

    def read_sheet(self, sheet: str) -> SheetData | None:
        try:
            data = self.client.request(
                "GET",
                f"{self._sheet_url(sheet)}/usedRange",
                params={"$select": "address,values"},
                headers=self._headers(),
            )
        except GraphError as exc:
            if "ItemNotFound" in str(exc) or "404" in str(exc):
                return None
            raise
        top_left = data["address"].split("!")[-1].split(":")[0].replace("$", "")
        m = re.match(r"([A-Z]+)(\d+)", top_left)
        col = 0
        for ch in m.group(1):
            col = col * 26 + (ord(ch) - 64)
        return SheetData(values=data["values"], top_row=int(m.group(2)), left_col=col)

    def write_cells(self, changes: list[CellChange]) -> None:
        for c in changes:
            self.client.request(
                "PATCH",
                f"{self._sheet_url(c.sheet)}/range(address='{_a1(c.row, c.col)}')",
                json={"values": [[c.new_value]]},
                headers=self._headers(),
            )

    def append_row(self, sheet: str, header: list[str], row: list[Any]) -> None:
        data = self.read_sheet(sheet)
        if data is None:
            return
        blank = all(str(v).strip() == "" for r in data.values for v in r)
        writes: list[tuple[int, list[Any]]] = []
        if blank:
            writes.append((1, header))
            next_row = 2
        else:
            next_row = data.top_row + len(data.values)
        writes.append((next_row, row))
        for r, vals in writes:
            addr = f"{_a1(r, 1)}:{_a1(r, len(vals))}"
            self.client.request(
                "PATCH",
                f"{self._sheet_url(sheet)}/range(address='{addr}')",
                json={"values": [[str(v) for v in vals]]},
                headers=self._headers(),
            )

    def close(self) -> None:
        if self._session_id:
            try:
                self.client.request(
                    "POST", f"{self.base}/closeSession", headers=self._headers(), expect_json=False
                )
            finally:
                self._session_id = None


class LocalWorkbook:
    def __init__(self, settings: Settings):
        self.path = settings.local_workbook_path
        self.wb = load_workbook(self.path)
        self._dirty = False

    def read_sheet(self, sheet: str) -> SheetData | None:
        if sheet not in self.wb.sheetnames:
            return None
        ws = self.wb[sheet]
        values = [
            ["" if v is None else v for v in row]
            for row in ws.iter_rows(min_row=ws.min_row, min_col=ws.min_column, values_only=True)
        ]
        return SheetData(values=values, top_row=ws.min_row, left_col=ws.min_column)

    def write_cells(self, changes: list[CellChange]) -> None:
        for c in changes:
            self.wb[c.sheet].cell(row=c.row, column=c.col, value=c.new_value)
        self._dirty = True

    def append_row(self, sheet: str, header: list[str], row: list[Any]) -> None:
        if sheet not in self.wb.sheetnames:
            return
        ws = self.wb[sheet]
        if ws.max_row == 1 and all(c.value in (None, "") for c in ws[1]):
            for i, h in enumerate(header, start=1):
                ws.cell(row=1, column=i, value=h)
        ws.append(row)
        self._dirty = True

    def close(self) -> None:
        if self._dirty:
            self.wb.save(self.path)
            self._dirty = False


def normalize_name(name: Any) -> str:
    return re.sub(r"\s+", " ", str(name or "")).strip().lower()


def to_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if 20000 < value < 80000:  # Excel serial date (Graph returns raw values)
            return EXCEL_EPOCH + timedelta(days=int(value))
        return None
    try:
        return date_parser.parse(str(value), dayfirst=False, fuzzy=False).date()
    except (ValueError, OverflowError):
        return None


@dataclass
class RosterRow:
    sheet_row: int
    date: date
    duty: str
    assignee: str


@dataclass
class Roster:
    sheet: str
    name_col: int  # 1-based worksheet column
    rows: list[RosterRow]
    staff_emails: dict[str, str] = field(default_factory=dict)  # normalized name -> email

    @property
    def people(self) -> list[str]:
        seen: dict[str, str] = {}
        for r in self.rows:
            if r.assignee:
                seen.setdefault(normalize_name(r.assignee), r.assignee)
        for name in self.staff_emails:
            seen.setdefault(name, name.title())
        return sorted(seen.values())

    @property
    def duties(self) -> list[str]:
        return sorted({r.duty for r in self.rows if r.duty})

    def email_for(self, name: str) -> str | None:
        return self.staff_emails.get(normalize_name(name))

    def name_for_email(self, email: str) -> str | None:
        email = email.lower()
        for name, addr in self.staff_emails.items():
            if addr.lower() == email:
                return name
        return None

    def find(self, person: str, day: date, duty: str | None) -> RosterRow:
        who = normalize_name(person)
        matches = [r for r in self.rows if r.date == day and normalize_name(r.assignee) == who]
        if duty:
            narrowed = [r for r in matches if normalize_name(duty) in normalize_name(r.duty)]
            matches = narrowed or matches
        if not matches:
            raise SwapError(f"{person} is not on the roster for {day:%a %d %b %Y}.")
        if len(matches) > 1:
            duties = ", ".join(r.duty or "?" for r in matches)
            raise SwapError(f"{person} has several duties on {day:%a %d %b %Y} ({duties}); please say which one.")
        return matches[0]

    def is_known(self, person: str) -> bool:
        return normalize_name(person) in {normalize_name(p) for p in self.people}


def _find_header(values: list[list[Any]], wanted: str) -> int:
    for i, row in enumerate(values[:15]):
        if any(normalize_name(v) == normalize_name(wanted) for v in row):
            return i
    raise SwapError(f"Could not find a '{wanted}' header in the first 15 rows of the sheet.")


def _col_index(header: list[Any], wanted: str, required: bool = True) -> int | None:
    for i, v in enumerate(header):
        if normalize_name(v) == normalize_name(wanted):
            return i
    if required:
        raise SwapError(f"Column '{wanted}' not found in header {header}.")
    return None


def load_roster(workbook: Workbook, settings: Settings) -> Roster:
    data = workbook.read_sheet(settings.roster_sheet)
    if data is None:
        raise SwapError(f"Worksheet '{settings.roster_sheet}' not found in the roster workbook.")
    h = _find_header(data.values, settings.date_column)
    header = data.values[h]
    di = _col_index(header, settings.date_column)
    ni = _col_index(header, settings.name_column)
    ui = _col_index(header, settings.duty_column, required=False) if settings.duty_column else None
    rows: list[RosterRow] = []
    for offset, raw in enumerate(data.values[h + 1 :], start=h + 1):
        day = to_date(raw[di] if di < len(raw) else None)
        if day is None:
            continue
        rows.append(
            RosterRow(
                sheet_row=data.top_row + offset,
                date=day,
                duty=str(raw[ui]).strip() if ui is not None and ui < len(raw) else "",
                assignee=str(raw[ni]).strip() if ni < len(raw) else "",
            )
        )

    staff: dict[str, str] = {}
    if settings.staff_sheet:
        sdata = workbook.read_sheet(settings.staff_sheet)
        if sdata is not None and sdata.values:
            try:
                sh = _find_header(sdata.values, settings.staff_name_column)
                sheader = sdata.values[sh]
                si = _col_index(sheader, settings.staff_name_column)
                ei = _col_index(sheader, settings.staff_email_column)
                for raw in sdata.values[sh + 1 :]:
                    name, email = (raw[si] if si < len(raw) else ""), (raw[ei] if ei < len(raw) else "")
                    if str(name).strip() and "@" in str(email):
                        staff[normalize_name(name)] = str(email).strip()
            except SwapError:
                pass

    return Roster(sheet=settings.roster_sheet, name_col=data.left_col + ni, rows=rows, staff_emails=staff)


def plan_changes(roster: Roster, req: SwapRequest, today: date, allow_past: bool) -> tuple[list[CellChange], list[str]]:
    """Validate a request against the roster and return the cell edits it needs, plus non-blocking warnings."""
    if not req.person_a or not req.person_b or not req.date_a:
        raise SwapError("The request is missing who is swapping and/or which date.")
    if normalize_name(req.person_a) == normalize_name(req.person_b):
        raise SwapError("Both sides of the swap are the same person.")
    if req.request_type == "swap" and not req.date_b:
        req = req.model_copy(update={"date_b": req.date_a})

    dates = [req.date_a] + ([req.date_b] if req.request_type == "swap" else [])
    if not allow_past and any(d < today for d in dates):
        raise SwapError("Swaps for dates that have already passed are not processed automatically.")

    row_a = roster.find(req.person_a, req.date_a, req.duty_a)
    warnings: list[str] = []

    if req.request_type == "cover":
        if not roster.is_known(req.person_b):
            raise SwapError(f"{req.person_b} is not on the roster or staff list.")
        canonical_b = next(p for p in roster.people if normalize_name(p) == normalize_name(req.person_b))
        clash = [r for r in roster.rows if r.date == row_a.date and normalize_name(r.assignee) == normalize_name(canonical_b)]
        if clash:
            warnings.append(f"{canonical_b} already has {clash[0].duty or 'a duty'} on {row_a.date:%a %d %b}.")
        return [
            CellChange(sheet=roster.sheet, row=row_a.sheet_row, col=roster.name_col,
                       old_value=row_a.assignee, new_value=canonical_b)
        ], warnings

    row_b = roster.find(req.person_b, req.date_b, req.duty_b)
    if row_a.sheet_row == row_b.sheet_row:
        raise SwapError("Both sides of the swap point at the same roster entry.")
    for person, day, own_row in ((row_b.assignee, row_a.date, row_b), (row_a.assignee, row_b.date, row_a)):
        clash = [r for r in roster.rows if r.date == day and normalize_name(r.assignee) == normalize_name(person)
                 and r.sheet_row not in (row_a.sheet_row, row_b.sheet_row)]
        if clash:
            warnings.append(f"{person} will also still have {clash[0].duty or 'another duty'} on {day:%a %d %b}.")
    return [
        CellChange(sheet=roster.sheet, row=row_a.sheet_row, col=roster.name_col,
                   old_value=row_a.assignee, new_value=row_b.assignee),
        CellChange(sheet=roster.sheet, row=row_b.sheet_row, col=roster.name_col,
                   old_value=row_b.assignee, new_value=row_a.assignee),
    ], warnings


def apply_to_roster(roster: Roster, changes: list[CellChange]) -> None:
    """Keep the in-memory roster in sync so later requests in the same run see earlier swaps."""
    by_row = {r.sheet_row: r for r in roster.rows}
    for c in changes:
        if c.row in by_row:
            by_row[c.row].assignee = c.new_value
