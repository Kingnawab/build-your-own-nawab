"""Create a demo roster workbook and inbox so the flow can be tried with DUTY_SWAP_BACKEND=local."""
from __future__ import annotations

import json
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

STAFF = [
    ("Alice Nguyen", "alice@example.com"),
    ("Bob Martinez", "bob@example.com"),
    ("Chen Wei", "chen@example.com"),
    ("Dana Okafor", "dana@example.com"),
]
DUTIES = ["Morning Gate", "Lunch Hall"]


def build(root: Path, start: date) -> None:
    root.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append(["Duty Roster"])
    ws.append(["Date", "Day", "Duty", "Assigned"])
    for i in range(14):
        day = start + timedelta(days=i)
        if day.weekday() >= 5:
            continue
        for j, duty in enumerate(DUTIES):
            ws.append([day, day.strftime("%A"), duty, STAFF[(i + j) % len(STAFF)][0]])
    for row in ws.iter_rows(min_row=3, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    staff = wb.create_sheet("Staff")
    staff.append(["Name", "Email"])
    for name, email in STAFF:
        staff.append([name, email])
    log = wb.create_sheet("Swap Log")
    log.append(["Processed At", "Requested By", "Change", "Cells", "Email Subject"])
    wb.save(root / "duty_roster.xlsx")

    inbox = root / "mail" / "inbox"
    inbox.mkdir(parents=True, exist_ok=True)
    emails = [
        {
            "id": "msg-001",
            "subject": "Duty swap request",
            "sender_name": "Alice Nguyen",
            "sender_email": "alice@example.com",
            "cc": ["bob@example.com"],
            "body": "Hi, Bob and I have agreed to swap duties: I'll take his next Tuesday and he takes my next Wednesday. Thanks, Alice",
        },
        {
            "id": "msg-002",
            "subject": "Team lunch Friday",
            "sender_name": "Chen Wei",
            "sender_email": "chen@example.com",
            "body": "Reminder that team lunch is at 12:30 on Friday.",
        },
    ]
    for e in emails:
        (inbox / f"{e['id']}.json").write_text(json.dumps(e, indent=2))
    print(f"Wrote {root / 'duty_roster.xlsx'} and {len(emails)} emails in {inbox}")


def main() -> None:
    root = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("sample_data")
    build(root, date.today())


if __name__ == "__main__":
    main()
